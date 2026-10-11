# -*- coding: utf-8 -*-
"""
Word 文档生成模块

排版规范（严格按要求）：
  · 标题：黑体、三号(16pt)、加粗、居中
  · 正文：仿宋、三号(16pt)、两端对齐、首行缩进 2 字符
  · 图片：随文保存并插入，居中、等比例缩放到页面宽度内，全部展示
"""
import os
from datetime import datetime

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TITLE_FONT = "黑体"
BODY_FONT = "仿宋"
LATIN_FONT = "Times New Roman"

# 文档底部署名
CREDIT = "设计：公歧子　微信：gongqizi0"

TITLE_SIZE = Pt(16)     # 三号
BODY_SIZE = Pt(16)      # 三号
META_SIZE = Pt(10.5)    # 五号（辅助信息）

PAGE_W = 21.0           # A4 宽 cm
MARGIN_LR = 2.5         # 左右页边距 cm
TEXT_WIDTH_CM = PAGE_W - 2 * MARGIN_LR   # 16cm 正文宽度
MAX_IMG_W = TEXT_WIDTH_CM - 0.4          # 图片最大宽度
MAX_IMG_H = 21.0                         # 图片最大高度 cm


# --------------------------------------------------------------------------
# 底层样式工具
# --------------------------------------------------------------------------
def _set_run_font(run, ea_font, latin_font=LATIN_FONT, size=None, bold=False, color=None):
    run.font.name = latin_font
    run.font.size = size or BODY_SIZE
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), latin_font)
    rFonts.set(qn("w:hAnsi"), latin_font)
    rFonts.set(qn("w:eastAsia"), ea_font)


def _set_first_line_chars(paragraph, chars=2):
    """首行缩进 N 个字符（Word 的 firstLineChars，单位 1/100 字符）"""
    pPr = paragraph._p.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    ind.set(qn("w:firstLineChars"), str(int(chars * 100)))
    ind.set(qn("w:firstLine"), str(int(chars * 16 * 20)))  # 兜底：2*16pt -> twips


def _set_line_spacing(paragraph, multiple=1.5):
    pf = paragraph.paragraph_format
    pf.line_spacing = multiple


def setup_document():
    doc = Document()
    # 页面设置
    sec = doc.sections[0]
    sec.page_width = Cm(PAGE_W)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(MARGIN_LR)
    sec.right_margin = Cm(MARGIN_LR)
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)

    # Normal 样式默认
    normal = doc.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = BODY_SIZE
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    return doc


def add_title(doc, text):
    """标题：黑体 三号 加粗 居中"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(6)
    pf.space_after = Pt(12)
    run = p.add_run(text)
    _set_run_font(run, TITLE_FONT, TITLE_FONT, TITLE_SIZE, bold=True)
    return p


def add_meta(doc, text):
    """辅助信息（来源/时间/链接）：仿宋 五号 居中 灰色"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    _set_run_font(run, BODY_FONT, LATIN_FONT, META_SIZE, bold=False,
                  color=RGBColor(0x66, 0x66, 0x66))
    return p


def add_body(doc, text):
    """正文：仿宋 三号 两端对齐 首行缩进 2 字符"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _set_first_line_chars(p, 2)
    _set_line_spacing(p, 1.5)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    _set_run_font(run, BODY_FONT, LATIN_FONT, BODY_SIZE, bold=False)
    return p


def add_caption(doc, text):
    """图注：仿宋 五号 居中"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    _set_run_font(run, BODY_FONT, LATIN_FONT, META_SIZE,
                  color=RGBColor(0x88, 0x88, 0x88))
    return p


def add_separator(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("— · — · — · — · — · — · — · — · —")
    _set_run_font(run, BODY_FONT, LATIN_FONT, Pt(10.5),
                  color=RGBColor(0xbb, 0xbb, 0xbb))
    return p


def add_image(doc, img_path):
    """插入图片：居中，限制宽高，保持比例，全部显示"""
    if not img_path or not os.path.exists(img_path):
        return None
    w_cm, h_cm = _image_size_cm(img_path)
    ratio = (h_cm / w_cm) if w_cm else 1
    width = MAX_IMG_W
    if width * ratio > MAX_IMG_H:
        width = MAX_IMG_H / ratio
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(img_path, width=Cm(width))
    return p


def _image_size_cm(path):
    try:
        from PIL import Image
        with Image.open(path) as im:
            w, h = im.size
        dpi = 96.0
        return w / dpi * 2.54, h / dpi * 2.54
    except Exception:
        return 10.0, 6.0


def add_page_break(doc):
    doc.add_page_break()


def add_credit(doc, text=None):
    """文档底部署名：黑体小四 居中"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text or CREDIT)
    _set_run_font(run, TITLE_FONT, TITLE_FONT, Pt(12), bold=True,
                  color=RGBColor(0x33, 0x33, 0x33))
    return p


# --------------------------------------------------------------------------
# 组装：单个品种 -> 章节内容（在已存在的 doc 上追加）
# --------------------------------------------------------------------------
def _fmt_date(s):
    s = (s or "").strip()
    return s


def render_variety_section(doc, variety, articles, img_dir=None, image_map=None,
                           year_label=""):
    """把单个品种的资讯渲染进 doc。
    articles: [{title,date,source,url,paragraphs,images,blocks,...}]
    image_map: {url: local_filename} 已下载的图片
    说明：正文与配图按原文顺序（blocks）混排，图片落在对应的资讯正文中。
    """
    add_title(doc, f"{variety['name']}（{variety['code']}）资讯汇编")
    ex_name = variety.get("ex_name", "")
    add_meta(doc, f"{ex_name} · 共 {len(articles)} 条 · 整理时间 "
                  f"{datetime.now().strftime('%Y-%m-%d %H:%M')}")

    if not articles:
        add_body(doc, "该品种在所选日期范围内暂无相关资讯。")
        return

    for i, art in enumerate(articles, 1):
        # 资讯标题：黑体 三号 加粗 居中
        add_title(doc, f"{i}. {art['title']}")
        src_label = art.get("src_label") or art.get("src") or "东方财富"
        meta = "　".join(x for x in [
            f"来源：{art.get('source') or src_label}",
            f"时间：{_fmt_date(art.get('date'))}",
        ] if x)
        add_meta(doc, meta)
        if art.get("url"):
            add_meta(doc, f"原文链接：{art['url']}")

        # 图文按原顺序混排
        blocks = art.get("blocks") or []
        if not blocks:
            blocks = [{"type": "p", "text": t} for t in (art.get("paragraphs") or [])]
            blocks += [{"type": "img", "url": u} for u in (art.get("images") or [])]
        if not any(b.get("type") == "p" for b in blocks) and art.get("summary"):
            blocks.insert(0, {"type": "p", "text": art["summary"]})

        for b in blocks:
            if b.get("type") == "img":
                local = (image_map or {}).get(b.get("url"))
                if local and img_dir:
                    fp = os.path.join(img_dir, local)
                    add_image(doc, fp)
                    add_caption(doc, "（原文配图）")
            else:
                txt = (b.get("text") or "").strip()
                if txt:
                    add_body(doc, txt)
        if i < len(articles):
            add_separator(doc)


def build_doc(path, title, blocks, image_map=None):
    """通用构建。
    blocks: [('variety', variety, articles), ...] 或 ('cover', text, None)
    """
    doc = setup_document()
    add_title(doc, title)
    add_meta(doc, f"资讯来源：东方财富　|　整理时间："
                  f"{datetime.now().strftime('%Y-%m-%d %H:%M')}")
    for kind, payload, extra in blocks:
        if kind == "variety":
            add_page_break(doc)
            render_variety_section(doc, payload, extra,
                                   img_dir=os.path.dirname(path),
                                   image_map=image_map)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    doc.save(path)
    return path
